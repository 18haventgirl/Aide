"""Generate real synthetic upload files; requires Pillow, python-docx, reportlab, pypdf.

Usage: python backend/tests/create_medical_parser_fixtures.py logs/parser-format-matrix
No patient information, credentials or network calls are used.
"""
import argparse
import json
from pathlib import Path

LINES = [
    '【虚构测试病历，不是真实患者资料】',
    '示例医院 门诊病历',
    '姓名：周示例', '性别：女', '年龄：38岁', '科室：全科门诊',
    '就诊日期：2026-10-09',
    '主诉：咳嗽3天。',
    '现病史：干咳3天，否认胸痛，无发热。',
    '既往史：否认高血压病史。',
    '过敏史：青霉素过敏。',
    '查体：体温36.7℃。',
    '初步诊断：上呼吸道感染（待排）。',
    '处理：本次未开药。',
    '医嘱：按原就诊医生安排复诊。',
]


def generate(directory: Path):
    from PIL import Image, ImageDraw, ImageFont
    from docx import Document
    from docx.shared import Pt
    from pypdf import PdfReader, PdfWriter
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont

    directory.mkdir(parents=True, exist_ok=True)
    prefix = 'matrix-20261009'
    text = '\n'.join(LINES)
    (directory / f'{prefix}.txt').write_text(text, encoding='utf-8-sig')

    document = Document()
    document.styles['Normal'].font.size = Pt(11)
    document.add_heading(LINES[0], level=1)
    document.add_paragraph(LINES[1])
    table = document.add_table(rows=0, cols=2)
    for line in LINES[2:7]:
        cells = table.add_row().cells
        key, value = line.split('：', 1)
        cells[0].text = key + '：'
        cells[1].text = value
    for line in LINES[7:]:
        document.add_paragraph(line)
    document.save(directory / f'{prefix}.docx')

    pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
    native = directory / f'{prefix}-native.pdf'
    pdf = canvas.Canvas(str(native), pagesize=(595, 842))
    pdf.setFont('STSong-Light', 14)
    for index, line in enumerate(LINES):
        pdf.drawString(40, 795 - index * 42, line)
    pdf.save()
    assert '周示例' in PdfReader(native).pages[0].extract_text()

    font_paths = ['C:/Windows/Fonts/msyh.ttc', '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc']
    font_path = next((path for path in font_paths if Path(path).exists()), None)
    if not font_path:
        raise RuntimeError('A CJK font is required to generate Chinese OCR fixtures.')
    image = Image.new('RGB', (1500, 1750), 'white')
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(font_path, 42)
    for index, line in enumerate(LINES):
        draw.text((70, 70 + index * 104), line, font=font, fill='black')
    for extension, format_name in [('png', 'PNG'), ('jpg', 'JPEG'), ('jpeg', 'JPEG')]:
        image.save(directory / f'{prefix}.{extension}', format=format_name, quality=94)
    scanned = directory / f'{prefix}-scanned.pdf'
    image.save(scanned, format='PDF', resolution=150)

    # One PDF with a native text page followed by an image-only page.
    writer = PdfWriter()
    writer.append(native)
    writer.append(scanned)
    writer.write(directory / f'{prefix}-mixed.pdf')
    image.close()
    manifest = []
    for suffix, methods in [('.txt', ['text']), ('.docx', ['docx']), ('-native.pdf', ['pdf']),
                            ('.png', ['ocr']), ('.jpg', ['ocr']), ('.jpeg', ['ocr']),
                            ('-scanned.pdf', ['ocr']), ('-mixed.pdf', ['pdf', 'ocr'])]:
        manifest.append({'filename': prefix + suffix, 'methods': methods,
                         'name': '周示例', 'date': '2026-10-09', 'diagnosis_contains': '待排',
                         'allergy_contains': '青霉素', 'corrected_name': '周示例（已核对）'})
    (directory / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f'Generated {len(manifest)} synthetic files in {directory.resolve()}')


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('directory', type=Path)
    generate(parser.parse_args().directory)
